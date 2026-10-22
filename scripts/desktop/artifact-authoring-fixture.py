"""Exercise bundled authoring libraries; outputs contain generated test data only."""
import hashlib,json,sys
from pathlib import Path

libraries,out,marker=Path(sys.argv[1]),Path(sys.argv[2]),sys.argv[3]
sys.path.insert(0,str(libraries.resolve()));out.mkdir(parents=True,exist_ok=False)
from PIL import Image,ImageDraw
from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook
import xlsxwriter
from pptx import Presentation
from pptx.util import Inches as SlideInches
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,Image as PdfImage,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab
from pypdf import PdfReader

image=Image.new('RGB',(640,360),'#18212f');drawing=ImageDraw.Draw(image);drawing.rectangle((30,35,610,315),fill='#377dff');drawing.text((50,60),marker,fill='white');image.save(out/'preview.png')
with Image.open(out/'preview.png') as saved:assert saved.size==(640,360) and saved.getpixel((40,40))==(55,125,255)

document=Document();document.add_heading('Agentopia authoring verification',0);document.add_paragraph(marker);document.add_paragraph('Résumé — generated fixture with a table and image.')
table=document.add_table(rows=1,cols=2);table.style='Table Grid';table.rows[0].cells[0].text='Item';table.rows[0].cells[1].text='Amount'
for label,value in [('Design','10'),('Build','20'),('Total','30')]:
    row=table.add_row().cells;row[0].text=label;row[1].text=value
document.add_picture(str(out/'preview.png'),width=Inches(5));document.sections[0].footer.paragraphs[0].text='Generated verification output';document.save(out/'Résumé verification.docx')
saved_document=Document(out/'Résumé verification.docx');assert marker in '\n'.join(p.text for p in saved_document.paragraphs) and len(saved_document.inline_shapes)==1 and saved_document.tables[0].rows[-1].cells[1].text=='30'

workbook=xlsxwriter.Workbook(out/'Estimates.xlsx');sheet=workbook.add_worksheet('Estimate');bold=workbook.add_format({'bold':True});money=workbook.add_format({'num_format':'$0.00'})
sheet.write_row(0,0,['Item','Amount'],bold);sheet.write_row(1,0,['Design',10]);sheet.write_row(2,0,['Build',20]);sheet.write(3,0,'Total',bold);sheet.write_formula(3,1,'=SUM(B2:B3)',money,30);sheet.write(5,0,marker);sheet.set_column(0,0,28);sheet.set_column(1,1,14);sheet.insert_image('D2',str(out/'preview.png'),{'x_scale':0.3,'y_scale':0.3})
chart=workbook.add_chart({'type':'column'});chart.add_series({'categories':'=Estimate!$A$2:$A$3','values':'=Estimate!$B$2:$B$3'});sheet.insert_chart('D10',chart);workbook.close()
formula=load_workbook(out/'Estimates.xlsx',data_only=False);values=load_workbook(out/'Estimates.xlsx',data_only=True);assert formula['Estimate']['B4'].value=='=SUM(B2:B3)' and values['Estimate']['B4'].value==30 and values['Estimate']['A6'].value==marker;formula.close();values.close()

presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[0]);slide.shapes.title.text='Agentopia authoring verification';slide.placeholders[1].text=marker;slide.notes_slide.notes_text_frame.text='Generated speaker notes '+marker
slide=presentation.slides.add_slide(presentation.slide_layouts[5]);slide.shapes.title.text='Saved results';slide.shapes.add_picture(str(out/'preview.png'),SlideInches(0.5),SlideInches(1.4),width=SlideInches(5));cells=slide.shapes.add_table(2,2,SlideInches(6),SlideInches(1.4),SlideInches(3),SlideInches(1)).table;cells.cell(0,0).text='Total';cells.cell(0,1).text='30';cells.cell(1,0).text='Status';cells.cell(1,1).text='Verified';presentation.save(out/'Results.pptx')
saved_presentation=Presentation(out/'Results.pptx');assert len(saved_presentation.slides)==2 and marker in saved_presentation.slides[0].notes_slide.notes_text_frame.text

font=Path(reportlab.__file__).parent/'fonts/Vera.ttf';assert font.is_file();pdfmetrics.registerFont(TTFont('VerificationVera',str(font)));styles=getSampleStyleSheet();styles['Normal'].fontName='VerificationVera';styles['Heading1'].fontName='VerificationVera'
pdf=SimpleDocTemplate(str(out/'Summary.pdf'));pdf.build([Paragraph('Agentopia authoring verification',styles['Heading1']),Paragraph(marker,styles['Normal']),Paragraph('Résumé verification',styles['Normal']),Spacer(1,18),Table([['Item','Amount'],['Design','10'],['Build','20'],['Total','30']]),Spacer(1,18),PdfImage(str(out/'preview.png'),width=320,height=180),PageBreak(),Paragraph('Second page '+marker,styles['Normal'])])
reader=PdfReader(out/'Summary.pdf');text='\n'.join(page.extract_text() for page in reader.pages);assert len(reader.pages)==2 and marker in text and 'Résumé' in text and '30' in text
files=[{'name':file.name,'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()} for file in sorted(out.iterdir())]
result={'marker':marker,'python':sys.version,'files':files,'checks':['PNG dimensions and pixels','Word paragraphs, table and embedded image','Spreadsheet formula and cached value read with a different library','Presentation slides and speaker notes','Two-page PDF text, Unicode and totals read with a different library'],'visualLayoutVerified':False,'providerRequests':0}
(out/'creation-evidence.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf8');print(json.dumps({'success':True,'files':len(files),'checks':result['checks']}))
