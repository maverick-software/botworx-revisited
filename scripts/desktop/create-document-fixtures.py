"""Developer-only fixture generator. Runtime document reading has no Python dependency."""
from pathlib import Path
from datetime import datetime
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
import zipfile
import xml.etree.ElementTree as ET
import argparse
import json
import secrets

arguments = argparse.ArgumentParser()
arguments.add_argument('--output', type=Path)
arguments.add_argument('--live', action='store_true')
options = arguments.parse_args()
root = options.output or Path(__file__).resolve().parents[2] / 'desktop/tests/fixtures/documents'
root.mkdir(parents=True, exist_ok=True)
codes = {key: (key + '-' + secrets.token_hex(4).upper()) if options.live else value for key, value in {
    'PDF': 'PDF-BRAVO-285', 'PDF2': 'PDF-PAGE-TWO-909', 'WORD': 'WORD-ALPHA-641',
    'SHEET': 'SHEET-CHARLIE-837', 'SLIDE': 'SLIDE-DELTA-572', 'SLIDE2': 'SLIDE-SECOND-147',
}.items()}

c = canvas.Canvas(str(root / 'sample.pdf'))
c.drawString(72, 740, codes['PDF'])
c.drawString(72, 710, 'Revenue: 1250.50')
c.showPage()
c.drawString(72, 740, codes['PDF2'])
c.save()
c = canvas.Canvas(str(root / 'image-only.pdf'))
c.rect(72, 600, 200, 100, fill=1)
c.save()
encrypted = PdfWriter()
for page in PdfReader(str(root / 'sample.pdf')).pages:
    encrypted.add_page(page)
encrypted.encrypt('fixture-password')
with (root / 'password.pdf').open('wb') as stream:
    encrypted.write(stream)

doc = Document()
doc.add_heading(codes['WORD'], 0)
doc.add_paragraph('Project notes: café and 東京. Approved total: 83.')
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = 'Product'
table.rows[0].cells[1].text = 'Units'
for product, units in [('Orchids', '7'), ('Ferns', '3')]:
    row = table.add_row()
    row.cells[0].text, row.cells[1].text = product, units
doc.sections[0].header.paragraphs[0].text = 'HEADER-DOCX-194'
doc.save(root / 'sample.docx')

book = Workbook()
sheet = book.active
sheet.title = 'Quarterly'
sheet.append(['Product', 'Units', 'Ship date', 'Active', 'Formula'])
sheet.append(['Orchids', 7, datetime(2026, 10, 15), True, '=B2*2'])
sheet.append(['Ferns', 3, datetime(2026, 10, 16), False, '=B2+B3'])
notes = book.create_sheet('Notes & totals')
notes['A1'] = codes['SHEET']
notes['B2'] = 'café and 東京'
notes.sheet_state = 'hidden'
book.save(root / 'sample.xlsx')
# Keep one real cached formula result alongside a formula with no cached result.
with zipfile.ZipFile(root / 'sample.xlsx') as source:
    entries = {name: source.read(name) for name in source.namelist()}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tree = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
tree.find(".//s:c[@r='E3']/s:v", ns).text = '10'
entries['xl/worksheets/sheet1.xml'] = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
with zipfile.ZipFile(root / 'sample.xlsx', 'w', zipfile.ZIP_DEFLATED) as target:
    for name, data in entries.items():
        target.writestr(name, data)

slides = Presentation()
first = slides.slides.add_slide(slides.slide_layouts[1])
first.shapes.title.text = codes['SLIDE']
first.placeholders[1].text = 'Launch plan: café and 東京.'
table = first.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
for row, values in enumerate([['Status', 'Owner'], ['Ready', 'Team']]):
    for col, value in enumerate(values):
        table.cell(row, col).text = value
first.notes_slide.notes_text_frame.text = 'NOTES-ECHO-219'
second = slides.slides.add_slide(slides.slide_layouts[1])
second.shapes.title.text = codes['SLIDE2']
second.placeholders[1].text = 'Second physical slide, first in presentation order.'
slide_id = slides.slides._sldIdLst[0]
slides.slides._sldIdLst.remove(slide_id)
slides.slides._sldIdLst.append(slide_id)
slides.save(root / 'sample.pptx')
if options.live:
    (root / 'codes.json').write_text(json.dumps(codes), encoding='utf-8')
print('Generated PDF, Word, Excel, and PowerPoint interoperability fixtures.')
