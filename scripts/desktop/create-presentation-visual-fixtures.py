"""Generated local PPTX fixtures; no user documents or remote data."""
from pathlib import Path
import sys,zipfile,xml.etree.ElementTree as ET
sys.path.append(str(Path(__file__).resolve().parents[2]/'desktop/resources/artifact-python'))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from PIL import Image,ImageDraw
out=Path(sys.argv[1]);out.mkdir();image=Image.new('RGB',(200,200),'#20b486');draw=ImageDraw.Draw(image);draw.ellipse((30,30,170,170),fill='#ffffff');image.save(out/'embedded.png')
p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
def text(slide,x,y,w,h,value,size,color,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));run=box.text_frame.paragraphs[0].add_run();run.text=value;run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color);return box
s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('102A43')
text(s,.6,.45,12,1,'Presentation preview',32,'FFFFFF',True);text(s,.6,1.4,12,.8,'Layout, image, chart and table remain together.',20,'D7E7F5')
s.shapes.add_picture(str(out/'embedded.png'),Inches(.8),Inches(2.6),Inches(2),Inches(2))
data=CategoryChartData();data.categories=['Jan','Feb','Mar'];data.add_series('Visits',[12,24,18]);chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(3.5),Inches(2.3),Inches(5),Inches(3),data).chart;chart.has_legend=False
table=s.shapes.add_table(3,2,Inches(9),Inches(2.6),Inches(3.5),Inches(2)).table
for row,values in enumerate([['Team','Status'],['Design','Ready'],['Build','Review']]):
    for col,value in enumerate(values):table.cell(row,col).text=value
text(s,.6,6.5,12,.6,'VISUAL-PPTX-001',16,'FFFFFF')
s=p.slides.add_slide(p.slide_layouts[6]);text(s,.6,.7,12,1,'Second slide',36,'174B6F',True);text(s,.6,2.0,12,1,'Navigation preserves the original presentation order.',22,'222222');text(s,.6,4.0,12,1,'<img src="https://example.invalid/text" onerror="alert(1)">',16,'222222');p.save(out/'rich.pptx')
with zipfile.ZipFile(out/'rich.pptx') as source,zipfile.ZipFile(out/'external.pptx','w',zipfile.ZIP_DEFLATED) as target:
    for item in source.infolist():
        data=source.read(item.filename)
        if item.filename=='ppt/slides/_rels/slide1.xml.rels':
            xml=ET.fromstring(data)
            for node in xml:
                if node.attrib.get('Type','').endswith('/image'):node.set('Target','https://example.invalid/blocked-media.png');node.set('TargetMode','External')
            data=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        target.writestr(item.filename,data)
print(out)
