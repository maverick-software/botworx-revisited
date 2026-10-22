"""Generate local preview fixtures without user or network data."""
from pathlib import Path
import sys,zipfile,xml.etree.ElementTree as ET
sys.path.append(str(Path(__file__).resolve().parents[2]/'desktop/resources/artifact-python'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.styles.colors import Color
from datetime import datetime
out=Path(sys.argv[1]);out.mkdir()
w=Workbook();s=w.active;s.title='Quarterly sales'
s.merge_cells('A1:D1');s['A1']='Quarterly overview';s['A1'].font=Font(name='Arial',size=18,bold=True,color='FFFFFF');s['A1'].fill=PatternFill('solid',fgColor='17365D');s['A1'].alignment=Alignment(horizontal='center',vertical='center');s.row_dimensions[1].height=36
for cell,value in zip(['A3','B3','C3','D3'],['Product','Revenue','Share','Ship date']):
 s[cell]=value;s[cell].font=Font(name='Arial',bold=True,color='FFFFFF');s[cell].fill=PatternFill('solid',fgColor='4F81BD')
s.append(['Orchids',1234.5,0.125,datetime(2026,10,19)]);s['B4'].number_format='"$"#,##0.00';s['C4'].number_format='0.0%';s['D4'].number_format='yyyy-mm-dd'
s['A5']='<img src="https://example.invalid/pixel" onerror="alert(1)">';s['B5']='=B4*2';s['B5'].number_format='"$"#,##0.00';s['C5']='=B4/2';s['D5']=True
s['A6']='Wrapped text that stays within the original column';s['A6'].alignment=Alignment(wrap_text=True,vertical='center');s.row_dimensions[6].height=45;s['B6']=42;s['B6'].font=Font(italic=True,color=Color(theme=4,tint=.4));s['B6'].border=Border(bottom=Side(style='double',color='FF0000'))
s.column_dimensions['A'].width=26;s.column_dimensions['B'].width=18;s.column_dimensions['C'].width=12;s.column_dimensions['D'].width=18;s.row_dimensions[7].hidden=True;s.column_dimensions['F'].hidden=True;s['J2500']='Far row marker';s['B10001']='Outside bounded grid'
t=w.create_sheet('Notes & totals');t.sheet_state='hidden';t['A1']='Second sheet, separate selection';t['B2']=99
w.save(out/'styled.xlsx')
# Supply a stored result for one formula. The other intentionally has no cache.
file=out/'styled.xlsx'
with zipfile.ZipFile(file) as z: parts={n:z.read(n) for n in z.namelist()}
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};tree=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
for c in tree.findall('.//s:c',ns):
 if c.attrib.get('r')=='B5':c.find('s:v',ns).text='2469'
parts['xl/worksheets/sheet1.xml']=ET.tostring(tree)
with zipfile.ZipFile(file,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
