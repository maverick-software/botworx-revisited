import csv,hashlib,json,sys,zipfile
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation

root=Path(sys.argv[1]).resolve()
rows=list(csv.DictReader((root/'intake.csv').open(encoding='utf-8')))
amounts=[(Decimal(r['Units'])*Decimal(r['Unit Price'])).quantize(Decimal('.01')) for r in rows]
total=sum(amounts)
book=root/'artifact-workbook.xlsx'
deck=root/'artifact-deck.pptx'
formulas=load_workbook(book,data_only=False)
values=load_workbook(book,data_only=True)
assert formulas.sheetnames[:2]==['Items','Summary'],formulas.sheetnames
clean=lambda value:str(value).replace(' ','').replace('$','').upper()
for index,(row,amount) in enumerate(zip(rows,amounts),2):
    assert formulas['Items'][f'A{index}'].value==row['Item']
    assert Decimal(str(values['Items'][f'B{index}'].value))==Decimal(row['Units'])
    assert Decimal(str(values['Items'][f'C{index}'].value))==Decimal(row['Unit Price'])
    assert clean(formulas['Items'][f'D{index}'].value) in [f'=B{index}*C{index}',f'=ROUND(B{index}*C{index},2)']
    assert Decimal(str(values['Items'][f'D{index}'].value)).quantize(Decimal('.01'))==amount
assert clean(formulas['Items']['D6'].value)=='=SUM(D2:D4)'
assert clean(formulas['Summary']['B2'].value) in ["='ITEMS'!D6",'=ITEMS!D6']
assert Decimal(str(values['Items']['D6'].value)).quantize(Decimal('.01'))==total
assert Decimal(str(values['Summary']['B2'].value)).quantize(Decimal('.01'))==total
assert formulas['Items'].freeze_panes=='A2'
assert sum(len(sheet._charts) for sheet in formulas)>=1
assert not [(sheet.title,cell.coordinate) for sheet in formulas for line in sheet for cell in line if cell.data_type=='e']
ppt=Presentation(deck)
assert len(ppt.slides)==2
texts=['\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in ppt.slides]
assert 'Furniture intake' in texts[0]
assert 'Cost breakdown' in texts[1]
marker=(root/'README.md').read_text(encoding='utf-8').strip().splitlines()[0]
assert any(marker in text for text in texts)
tables=[shape.table for slide in ppt.slides for shape in slide.shapes if shape.has_table]
charts=[shape.chart for slide in ppt.slides for shape in slide.shapes if shape.has_chart]
assert tables and charts
table_text='\n'.join(cell.text for table in tables for row in table.rows for cell in row.cells)
assert all(row['Item'] in table_text for row in rows)
assert str(total) in table_text.replace(',','')
assert any([Decimal(str(v)).quantize(Decimal('.01')) for v in series.values]==amounts for chart in charts for series in chart.series)
external=[]
for file in [book,deck]:
    with zipfile.ZipFile(file) as package:
        assert package.testzip() is None
        external.extend(name for name in package.namelist() if name.endswith('.rels') and b'TargetMode="External"' in package.read(name))
assert not external,external
print(json.dumps({'success':True,'root':str(root),'total':str(total),'amounts':[str(v) for v in amounts],'formulaTotal':formulas['Items']['D6'].value,'summaryFormula':formulas['Summary']['B2'].value,'formulaCachesIndependentlyChecked':True,'generalOfficeRecalculationRun':False,'slides':2,'tables':len(tables),'charts':len(charts),'titles':['Furniture intake','Cost breakdown'],'marker':marker,'hashes':{file.name:hashlib.sha256(file.read_bytes()).hexdigest() for file in [book,deck]}}))
